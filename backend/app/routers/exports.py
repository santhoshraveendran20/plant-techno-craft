"""FR-26: export bills, payments, dues and stock to Excel / PDF for the accountant."""

import io
from datetime import date, datetime, timedelta
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..models import Bill, Payment, User
from ..security import require_owner
from ..services import analytics, app_settings
from ..services.audit import log_audit
from ..utils import today

router = APIRouter(prefix="/api/exports", tags=["exports"])

LOGO = Path(__file__).resolve().parent.parent / "assets" / "mark.png"
BRAND_DEEP = colors.HexColor("#124f2c")

STATUS_LABELS = {"part_paid": "Part-paid"}


def _range(date_from: date | None, date_to: date | None) -> tuple[date, date, datetime, datetime]:
    date_to = date_to or today()
    date_from = date_from or (date_to - timedelta(days=29))
    if date_from > date_to:
        raise HTTPException(400, "date_from must be on or before date_to")
    return (date_from, date_to, datetime.combine(date_from, datetime.min.time()),
            datetime.combine(date_to + timedelta(days=1), datetime.min.time()))


def _bills(db: Session, start: datetime, end: datetime):
    headers = ["Bill no", "Date", "Customer", "Mobile", "Status", "Subtotal", "Discount", "Total", "Returned",
               "Paid", "Due"]
    rows = []
    for b in db.execute(
        select(Bill).options(selectinload(Bill.customer))
        .where(Bill.bill_type == "sale", Bill.completed_at >= start, Bill.completed_at < end)
        .order_by(Bill.completed_at)
    ).scalars():
        rows.append([
            b.bill_no, b.completed_at.strftime("%Y-%m-%d %H:%M"), b.customer.name if b.customer else "Walk-in",
            b.customer.mobile if b.customer else "", STATUS_LABELS.get(b.status, b.status.title()),
            float(b.subtotal), float(b.discount), float(b.total), float(b.returned_amount), float(b.paid_amount),
            float(b.due_amount),
        ])
    return headers, rows, {5, 6, 7, 8, 9, 10}


def _payments(db: Session, start: datetime, end: datetime):
    headers = ["Date", "Type", "Mode", "Amount", "Bill no", "Customer", "UTR", "Payer UPI ID", "Payer name",
               "Status", "Match", "Source"]
    rows = []
    for p in db.execute(
        select(Payment).options(selectinload(Payment.bill), selectinload(Payment.customer))
        .where(Payment.created_at >= start, Payment.created_at < end).order_by(Payment.created_at)
    ).scalars():
        rows.append([
            p.created_at.strftime("%Y-%m-%d %H:%M"), p.kind, p.mode.upper(), float(p.amount),
            p.bill.bill_no if p.bill else "", p.customer.name if p.customer else "", p.utr or "",
            p.payer_vpa or "", p.payer_name or "", p.gateway_status, p.match_status, p.source,
        ])
    return headers, rows, {3}


def _dues(db: Session):
    from .customers import dues_list

    data = dues_list(q=None, min_days=0, min_amount=0, bucket=None, sort="amount", _=None, db=db)
    headers = ["Customer", "Mobile", "Amount due", "Oldest due date", "Days overdue", "Ageing", "Last reminder"]
    rows = [[r["name"], r["mobile"], r["amount_due"], r["oldest_due_date"] or "", r["days_overdue"], r["bucket"],
             (r["last_reminder_at"] or "")[:16].replace("T", " ")] for r in data["rows"]]
    return headers, rows, {2}


def _stock(db: Session):
    data = analytics.stock_report(db)
    headers = ["Product", "Category", "Barcode", "Unit", "Stock", "Reorder level", "Cost price", "Selling price",
               "Value at cost", "Value at selling", "Status"]
    rows = [[r["name"], r["category"] or "", r["barcode"] or "", r["unit"], r["stock_qty"], r["reorder_level"],
             r["cost_price"], r["selling_price"], r["cost_value"], r["selling_value"], r["status"]]
            for r in data["rows"]]
    return headers, rows, {6, 7, 8, 9}


def _sales(db: Session, date_from: date, date_to: date):
    data = analytics.sales_by_product(db, date_from, date_to)
    headers = ["Product", "Qty sold", "Revenue", "Cost", "Gross margin", "Margin %"]
    rows = [[r["name"], r["qty"], r["revenue"], r["cost"], r["margin"], r["margin_pct"]] for r in data["rows"]]
    return headers, rows, {2, 3, 4}


def _xlsx(title: str, headers: list[str], rows: list[list], money_cols: set[int]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="18181B")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for idx in money_cols:
        letter = get_column_letter(idx + 1)
        for cell in ws[letter][1:]:
            cell.number_format = "#,##0.00"
    for idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows[:500]]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 45)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(title: str, subtitle: str, headers: list[str], rows: list[list], money_cols: set[int], tagline: str = "") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)

    def fmt(value, idx):
        if idx in money_cols and isinstance(value, (int, float)):
            return f"{value:,.2f}"
        return Paragraph(str(value), cell_style)

    data = [headers] + [[fmt(v, i) for i, v in enumerate(r)] for r in rows]
    if not rows:
        data.append(["No records"] + [""] * (len(headers) - 1))
    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_DEEP),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d6dccf")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5ee")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for idx in money_cols:
        style.append(("ALIGN", (idx, 1), (idx, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    title_style = styles["Title"].clone("brandTitle", alignment=0, textColor=BRAND_DEEP, spaceAfter=2)
    text = [Paragraph(title, title_style)]
    if tagline:
        text.append(Paragraph(tagline.upper(), styles["Normal"].clone("tagline", fontSize=7.5, textColor=colors.HexColor("#578417"))))
    text.append(Paragraph(subtitle, styles["Normal"]))
    header = text
    if LOGO.is_file():
        header = [Table([[Image(str(LOGO), width=16 * mm, height=16 * mm, mask="auto"), text]], colWidths=[20 * mm, None],
                        style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)])]
    doc.build([*header, Spacer(1, 6 * mm), table])
    return buf.getvalue()


@router.get("/{kind}")
def export(
    kind: str,
    fmt: str = Query(default="xlsx", alias="format", pattern="^(xlsx|pdf)$"),
    date_from: date | None = None,
    date_to: date | None = None,
    owner: User = Depends(require_owner),
    db: Session = Depends(get_db),
):
    d_from, d_to, start, end = _range(date_from, date_to)
    period = f"{d_from:%d %b %Y} - {d_to:%d %b %Y}"
    if kind == "bills":
        headers, rows, money_cols = _bills(db, start, end)
        title = "Bills"
    elif kind == "payments":
        headers, rows, money_cols = _payments(db, start, end)
        title = "Payments"
    elif kind == "dues":
        headers, rows, money_cols = _dues(db)
        title, period = "Outstanding dues", f"As on {today():%d %b %Y}"
    elif kind == "stock":
        headers, rows, money_cols = _stock(db)
        title, period = "Stock report", f"As on {today():%d %b %Y}"
    elif kind == "sales":
        headers, rows, money_cols = _sales(db, d_from, d_to)
        title = "Sales by product"
    else:
        raise HTTPException(404, "Unknown export. Use bills, payments, dues, stock or sales.")
    shop = app_settings.get_value(db, "shop_name")
    tagline = app_settings.get_value(db, "shop_tagline")
    log_audit(db, owner, "export", kind, None, f"Exported {title.lower()} ({fmt}) {period}")
    db.commit()
    filename = f"{kind}_{d_from:%Y%m%d}_{d_to:%Y%m%d}.{fmt}" if kind in ("bills", "payments", "sales") else f"{kind}_{today():%Y%m%d}.{fmt}"
    if fmt == "xlsx":
        content = _xlsx(title, headers, rows, money_cols)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        content = _pdf(f"{shop} - {title}", period, headers, rows, money_cols, tagline)
        media = "application/pdf"
    return StreamingResponse(io.BytesIO(content), media_type=media,
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})
